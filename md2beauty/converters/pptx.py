from __future__ import annotations

from typing import Iterable, Optional, Union
from pptx import Presentation  # type: ignore
from pptx.util import Inches, Pt  # type: ignore
from pptx.dml.color import RGBColor  # type: ignore

from ..theme import Theme
from ..mdparser import (
    Document,
    Heading as IRHeading,
    Paragraph as IRParagraph,
    CodeBlock as IRCode,
    ListBlock as IRList,
    Image as IRImage,
    SvgBlock as IRSvg,
    ThematicBreak as IRHr,
    Table as IRTable,
)
from ..embeds import default_diagram_service
from . import Converter


class PptxConverter(Converter):
    def __init__(self, theme: Optional[Union[str, Theme]] = None):
        super().__init__(theme)

    def convert_document(self, doc: Document) -> bytes:  # type: ignore[override]
        prs = Presentation()
        if isinstance(self.theme, Theme):
            self.theme.apply_to_pptx(prs)

        slide = prs.slides.add_slide(prs.slide_layouts[5])
        top = Inches(1)
        left = Inches(1)
        width = Inches(8)
        height = Inches(5.5)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.clear()

        code_bg = self._code_background_color()
        if code_bg:
            try:
                fill = txBox.fill
                fill.solid()
                rgb = self._to_rgb(code_bg)
                if rgb:
                    fill.fore_color.rgb = rgb
            except Exception:
                pass

        for block in doc.blocks:
            if isinstance(block, IRHeading):
                p = tf.add_paragraph()
                p.level = 0
                p.text = block.text
                p.font.size = Pt(28)
                p.font.bold = True
            elif isinstance(block, IRParagraph):
                p = tf.add_paragraph()
                p.level = 0
                p.text = block.text
                p.font.size = Pt(18)
            elif isinstance(block, IRCode):
                if block.language == 'mermaid':
                    try:
                        rr = default_diagram_service.render(
                            kind="mermaid",
                            code=block.code,
                            preferred_formats=["png"],
                        )
                    except Exception:
                        rr = None
                    if rr and rr.path:
                        try:
                            slide.shapes.add_picture(rr.path, left, top)
                        finally:
                            rr.cleanup()
                    else:
                        p = tf.add_paragraph()
                        p.text = block.code
                        p.font.name = 'Consolas'
                        p.font.size = Pt(12)
                else:
                    lang = (block.language or '').strip()
                    text = block.code
                    if lang:
                        try:
                            from pygments import lex  # type: ignore
                            from pygments.lexers import get_lexer_by_name  # type: ignore
                            from pygments.styles import get_style_by_name  # type: ignore
                            lexer = get_lexer_by_name(lang)
                            style_name = None
                            if isinstance(self.theme, Theme):
                                ppt_fmt = self.theme.formats.get('pptx') if isinstance(self.theme.formats, dict) else None
                                if ppt_fmt and getattr(ppt_fmt, 'pygments_style', None):
                                    style_name = ppt_fmt.pygments_style  # type: ignore[attr-defined]
                                if not style_name and getattr(self.theme, 'pygments_style', None):
                                    style_name = self.theme.pygments_style  # type: ignore[attr-defined]
                            st = get_style_by_name(style_name) if style_name else None
                            color_map = {}
                            if st:
                                for ttype, ndef in st.styles.items():
                                    color_map[str(ttype)] = ndef
                            p = tf.add_paragraph()
                            p.level = 0
                            p.font.name = 'Consolas'
                            p.font.size = Pt(12)
                            for ttype, val in lex(text, lexer):
                                if not val:
                                    continue
                                run = p.add_run()
                                run.text = val
                                ndef = color_map.get(str(ttype), '')
                                if 'bold' in ndef:
                                    run.font.bold = True
                                if 'italic' in ndef or 'em' in ndef:
                                    run.font.italic = True
                                m = None
                                if '#' in ndef:
                                    try:
                                        m = ndef.split('#', 1)[1].split()[0]
                                    except Exception:
                                        m = None
                                rgb = self._to_rgb('#' + m) if m else None
                                if rgb:
                                    run.font.color.rgb = rgb
                        except Exception:
                            p = tf.add_paragraph()
                            p.text = text
                            p.font.name = 'Consolas'
                            p.font.size = Pt(12)
                    else:
                        p = tf.add_paragraph()
                        p.text = text
                        p.font.name = 'Consolas'
                        p.font.size = Pt(12)
            elif isinstance(block, IRList):
                from ..mdparser import parse_inline_emphasis as _pie
                for item in block.items:
                    p = tf.add_paragraph()
                    p.level = 1 if block.ordered else 1
                    p.text = item.text
                    p.font.size = Pt(18)
            elif isinstance(block, IRImage):
                try:
                    slide.shapes.add_picture(block.src, left, top)
                except Exception:
                    p = tf.add_paragraph()
                    p.text = f"[image: {block.alt}]"
            elif isinstance(block, IRSvg):
                p = tf.add_paragraph()
                p.text = "[svg not supported]"
            elif isinstance(block, IRHr):
                p = tf.add_paragraph()
                p.text = "—" * 20
            elif isinstance(block, IRTable):
                p = tf.add_paragraph()
                p.text = "[table omitted]"

        from io import BytesIO
        buf = BytesIO()
        prs.save(buf)
        return buf.getvalue()

    def _to_rgb(self, s: str) -> Optional[RGBColor]:
        try:
            v = s.lstrip('#')
            if len(v) == 3:
                v = ''.join(ch*2 for ch in v)
            if len(v) == 6:
                return RGBColor.from_string(v.upper())
        except Exception:
            return None
        return None

    def _code_background_color(self) -> Optional[str]:
        if isinstance(self.theme, Theme):
            try:
                st = None
                fmt = self.theme.formats.get('pptx') if isinstance(self.theme.formats, dict) else None
                if fmt and hasattr(fmt, 'styles'):
                    st = (fmt.styles.get('pre') or fmt.styles.get('code'))
                if not st:
                    st = (self.theme.styles.get('pre') or self.theme.styles.get('code'))
                if st and getattr(st, 'background_color', None):
                    return str(st.background_color)
            except Exception:
                return None
        return None
