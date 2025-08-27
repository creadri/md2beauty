from __future__ import annotations

from typing import Iterable, Optional, Union
from docx import Document  # type: ignore
from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
from docx.shared import Pt, RGBColor  # type: ignore
from docx.oxml.shared import OxmlElement, qn  # type: ignore

from ..theme import Theme
from .. import is_debug
from ..mdparser import parse_markdown_stream
from ..embeds import default_diagram_service
from . import Converter


class DocxConverter(Converter):
    def __init__(self, theme: Optional[Union[str, Theme]] = None):
        super().__init__(theme)

    def convert_stream(self, lines: Iterable[str]) -> bytes:
        document = Document()
        if isinstance(self.theme, Theme):
            self.theme.apply_to_docx(document)
        from ..mdparser import (
            Heading as IRHeading,
            Paragraph as IRParagraph,
            CodeBlock as IRCode,
            ListBlock as IRList,
            Image as IRImage,
            SvgBlock as IRSvg,
            ThematicBreak as IRHr,
            Table as IRTable,
            parse_inline_emphasis,
        )
        for block in parse_markdown_stream(lines=lines):
            if isinstance(block, IRHeading):
                lvl = max(1, min(6, block.level))
                par = document.add_paragraph()
                par.style = f'Heading {lvl}' if f'Heading {lvl}' in [s.name for s in document.styles] else par.style
                runs = parse_inline_emphasis(block.text)
                for span in runs:
                    r = par.add_run(span.text)
                    r.bold = bool(span.bold)
                    r.italic = bool(span.italic)
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif isinstance(block, IRParagraph):
                par = document.add_paragraph()
                runs = parse_inline_emphasis(block.text)
                for span in runs:
                    r = par.add_run(span.text)
                    r.bold = bool(span.bold)
                    r.italic = bool(span.italic)
            elif isinstance(block, IRCode):
                if block.language == 'mermaid':
                    try:
                        rr = default_diagram_service.render(
                            kind="mermaid",
                            code=block.code,
                            preferred_formats=["png"],
                        )
                    except Exception:
                        if is_debug():
                            raise
                        rr = None
                    if rr and rr.path:
                        try:
                            document.add_picture(rr.path)
                        finally:
                            rr.cleanup()
                    else:
                        par = document.add_paragraph()
                        par.add_run(block.code)
                else:
                    par = document.add_paragraph()
                    # background shading from theme
                    bg_hex = self._code_background_color()
                    if bg_hex:
                        try:
                            p = par._p
                            pPr = p.get_or_add_pPr()
                            shd = OxmlElement('w:shd')
                            shd.set(qn('w:val'), 'clear')
                            shd.set(qn('w:color'), 'auto')
                            shd.set(qn('w:fill'), bg_hex.lstrip('#'))
                            pPr.append(shd)
                        except Exception:
                            pass
                    lang = (block.language or '').strip()
                    text = block.code
                    highlighted = False
                    if lang:
                        try:
                            from pygments import lex  # type: ignore
                            from pygments.lexers import get_lexer_by_name  # type: ignore
                            from pygments.token import Token  # type: ignore
                            lexer = get_lexer_by_name(lang)
                            style_name = None
                            if isinstance(self.theme, Theme):
                                html_fmt = self.theme.formats.get('docx') if isinstance(self.theme.formats, dict) else None
                                if html_fmt and getattr(html_fmt, 'pygments_style', None):
                                    style_name = html_fmt.pygments_style  # type: ignore[attr-defined]
                                if not style_name and getattr(self.theme, 'pygments_style', None):
                                    style_name = self.theme.pygments_style  # type: ignore[attr-defined]
                            colors = {}
                            if style_name:
                                try:
                                    from pygments.styles import get_style_by_name  # type: ignore
                                    st = get_style_by_name(style_name)
                                    for ttype, ndef in st.styles.items():
                                        colors[str(ttype)] = ndef
                                except Exception:
                                    pass
                            for ttype, val in lex(text, lexer):
                                if val == '':
                                    continue
                                r = par.add_run(val)
                                # basic emphasis mapping via token type string contains
                                tname = str(ttype)
                                ndef = colors.get(tname, '')
                                if 'bold' in ndef:
                                    r.bold = True
                                if 'italic' in ndef or 'em' in ndef:
                                    r.italic = True
                                m = None
                                if '#' in ndef:
                                    try:
                                        m = ndef.split('#', 1)[1].split()[0]
                                    except Exception:
                                        m = None
                                if m and len(m) in (3, 6):
                                    try:
                                        if len(m) == 3:
                                            m = ''.join(ch*2 for ch in m)
                                        r.font.color.rgb = RGBColor.from_string(m.upper())
                                    except Exception:
                                        pass
                            highlighted = True
                        except Exception:
                            highlighted = False
                    if not highlighted:
                        r = par.add_run(text)
                        r.font.name = 'Consolas'
                        r.font.size = Pt(10)
            elif isinstance(block, IRList):
                from ..mdparser import parse_inline_emphasis as _pie
                for item in block.items:
                    par = document.add_paragraph(style='List Number' if block.ordered else 'List Bullet')
                    spans = _pie(item.text)
                    for span in spans:
                        r = par.add_run(span.text)
                        r.bold = bool(span.bold)
                        r.italic = bool(span.italic)
            elif isinstance(block, IRImage):
                try:
                    document.add_picture(block.src)
                except Exception:
                    par = document.add_paragraph()
                    par.add_run(f"[image: {block.alt}]")
            elif isinstance(block, IRSvg):
                # No native SVG support in docx; skipping
                par = document.add_paragraph()
                par.add_run("[svg not supported]")
            elif isinstance(block, IRHr):
                document.add_paragraph().add_run("—" * 20)
            elif isinstance(block, IRTable):
                rows = len(block.rows) + 1
                cols = max(1, len(block.headers)) if block.headers else len(block.rows[0]) if block.rows else 1
                table = document.add_table(rows=rows, cols=cols)
                # header
                for j, h in enumerate(block.headers):
                    table.cell(0, j).text = h
                # body
                for i, row in enumerate(block.rows, start=1):
                    for j, cell in enumerate(row):
                        table.cell(i, j).text = cell

        from io import BytesIO
        buf = BytesIO()
        document.save(buf)
        return buf.getvalue()

    def _code_background_color(self) -> Optional[str]:
        if isinstance(self.theme, Theme):
            try:
                st = None
                fmt = self.theme.formats.get('docx') if isinstance(self.theme.formats, dict) else None
                if fmt and hasattr(fmt, 'styles'):
                    st = (fmt.styles.get('pre') or fmt.styles.get('code'))
                if not st:
                    st = (self.theme.styles.get('pre') or self.theme.styles.get('code'))
                if st and getattr(st, 'background_color', None):
                    return str(st.background_color)
            except Exception:
                return None
        return None
