from md2beauty.theme import Theme
from md2beauty.converters import HtmlConverter


def test_apply_to_docx_runs(tmp_path):
    try:
        from docx import Document
    except Exception:
        return  # skip if python-docx unavailable
    theme = Theme.from_dict({
        "page": {"type": "a3", "orientation": "landscape", "margins": {"top": "1in", "right": "0.5in", "bottom": "1in", "left": "0.5in"}},
        "styles": {"h1": {"font-size": "24pt"}, "p": {"font-size": "12pt"}}
    })
    doc = Document()
    theme.apply_to_docx(doc)
    # If no exception, consider success; optionally save to ensure no serialization error
    out = tmp_path / "out.docx"
    doc.save(out)
    assert out.exists()


def test_html_includes_page_css():
    theme = Theme.from_dict({
        "page": {"type": "letter", "margins": {"top": "1in", "right": "1in", "bottom": "1in", "left": "1in"}},
        "styles": {"p": {"color": "#111"}}
    })
    html = HtmlConverter(theme).convert("# T\n\nPara").decode("utf-8")
    assert "@page" in html and "md2beauty-page" in html


def test_apply_to_pptx_runs_with_dummy():
    class DummyFont:
        def __init__(self):
            self.name = None
            self.size = None
            self.color = type("C", (), {"rgb": None})()
            self.bold = None
            self.italic = None
    class DummyRun:
        def __init__(self):
            self.font = DummyFont()
    class DummyParagraph:
        def __init__(self):
            self.runs = [DummyRun()]
    class DummyTextFrame:
        def __init__(self):
            self.paragraphs = [DummyParagraph()]
    class DummyShape:
        def __init__(self, is_title=False):
            self.is_title = is_title
            self.text_frame = DummyTextFrame()
    class DummyLayout:
        def __init__(self):
            self.shapes = [DummyShape(True), DummyShape(False)]
    class DummyPres:
        def __init__(self):
            self.slide_layouts = [DummyLayout(), DummyLayout()]
    theme = Theme.from_dict({"styles": {"h1": {"font-size": "32pt"}, "p": {"font-size": "18pt"}}})
    pres = DummyPres()
    theme.apply_to_pptx(pres)
    # If no exception, assume applied
    assert True
